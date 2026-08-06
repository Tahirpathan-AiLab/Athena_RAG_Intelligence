import re
import uuid
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.database import SessionLocal
from app.models import Document, DocumentChunk
from app.services.embeddings import get_embedding

UPLOAD_DIR = Path("data/uploads")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


def save_uploaded_file(content: bytes, original_filename: str | None = None) -> Path:
    original_name = Path(original_filename or "upload").name
    suffix = Path(original_name).suffix.lower() or ".txt"
    safe_name = f"{uuid.uuid4().hex}{suffix}"
    destination = UPLOAD_DIR / safe_name

    destination.write_bytes(content)
    return destination


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(file_path)
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return "\n".join(text_parts)

    if suffix == ".docx":
        document = DocxDocument(file_path)
        return "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        )

    if suffix in {".txt", ".md"}:
        return file_path.read_text(encoding="utf-8")

    raise ValueError(f"Unsupported file type: {suffix}")


def split_text_into_chunks(text: str, chunk_size: int = 350) -> list[str]:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return []

    words = cleaned.split()
    chunks = []
    for index in range(0, len(words), chunk_size):
        chunk = " ".join(words[index : index + chunk_size])
        chunks.append(chunk)
    return chunks


def ingest_document(
    file_path: Path, filename: str | None = None, category: str = "General"
):
    session = SessionLocal()

    try:
        text = extract_text(file_path)
        chunks = split_text_into_chunks(text)
        if not chunks:
            raise ValueError("No readable content found in the uploaded document")

        document = Document(
            filename=filename or file_path.name,
            relative_path=str(file_path),
            category=category,
        )
        session.add(document)
        session.flush()

        for index, chunk in enumerate(chunks):
            embedding = get_embedding(chunk)
            session.add(
                DocumentChunk(
                    document_id=document.id,
                    chunk_index=index,
                    content=chunk,
                    embedding=embedding,
                )
            )

        session.commit()
        session.refresh(document)
        return {
            "filename": document.filename,
            "chunks": len(chunks),
            "document_id": document.id,
        }
    except Exception as exc:
        session.rollback()
        raise exc
    finally:
        session.close()


def ingest_pdf(pdf_path: Path):
    return ingest_document(pdf_path, filename=pdf_path.name)
